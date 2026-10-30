"""#2139 — a zone's records as CSV / XLSX: export, template, import.

The contract: the export is one row per record with the structured columns
split out; re-importing an untouched export is a no-op; a row goes through the
record API's own checks (a bad one is a per-row error with its row number,
and the commit refuses the whole file); IPAM-owned records are never changed
or deleted; the CNAME rule skips a clashing row; formula injection is escaped
on export and a leading ``=`` is text on import; and the commit reaches the
servers as record ops on one serial bump, like the zone-file import.
"""

from __future__ import annotations

import csv
import io
import uuid
from typing import Any

from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.security import create_access_token, hash_password
from app.models.audit import AuditLog
from app.models.auth import Group, Role, User
from app.models.dns import DNSRecord, DNSRecordOp, DNSServer, DNSServerGroup, DNSView, DNSZone
from app.services.dns_io import spreadsheet as sheet_io

ZONE = "sheet.example.edu."
COLS = sheet_io.RECORD_COLUMNS


async def _headers(db: AsyncSession, perms: list[dict] | None = None) -> dict[str, str]:
    tag = uuid.uuid4().hex[:8]
    user = User(
        username=f"u-{tag}",
        email=f"{tag}@example.test",
        display_name="Test",
        hashed_password=hash_password("x" * 12),
        is_superadmin=perms is None,
    )
    if perms is not None:
        role = Role(name=f"r-{tag}", permissions=perms)
        group = Group(name=f"g-{tag}")
        group.roles = [role]
        user.groups = [group]
        db.add_all([role, group])
    db.add(user)
    await db.flush()
    return {"Authorization": f"Bearer {create_access_token(str(user.id))}"}


async def _zone(db: AsyncSession, servers: int = 1) -> tuple[DNSServerGroup, DNSZone, list]:
    grp = DNSServerGroup(name=f"g-{uuid.uuid4().hex[:6]}")
    db.add(grp)
    await db.flush()
    rows = [
        DNSServer(
            group_id=grp.id,
            driver="bind9",
            host=f"ns{i}.example.edu",
            name=f"srv-{uuid.uuid4().hex[:6]}",
            is_primary=i == 0,
            is_enabled=True,
        )
        for i in range(servers)
    ]
    db.add_all(rows)
    zone = DNSZone(
        group_id=grp.id,
        name=ZONE,
        zone_type="primary",
        kind="forward",
        primary_ns=f"ns1.{ZONE}",
        admin_email=f"admin.{ZONE}",
    )
    db.add(zone)
    await db.flush()
    return grp, zone, rows


def _rec(zone: DNSZone, name: str, rtype: str, value: str, **kw: Any) -> DNSRecord:
    return DNSRecord(
        zone_id=zone.id,
        name=name,
        fqdn=f"{name}.{zone.name}" if name != "@" else zone.name,
        record_type=rtype,
        value=value,
        **kw,
    )


def _csv(rows: list[dict[str, Any]]) -> bytes:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=COLS, extrasaction="ignore")
    w.writeheader()
    for r in rows:
        w.writerow(r)
    return buf.getvalue().encode()


def _read_csv(data: bytes) -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(data.decode())))


def _base(grp: DNSServerGroup, zone: DNSZone) -> str:
    return f"/api/v1/dns/groups/{grp.id}/zones/{zone.id}"


async def _preview(
    client: AsyncClient,
    h: dict[str, str],
    grp: DNSServerGroup,
    zone: DNSZone,
    data: bytes,
    *,
    mode: str = "merge",
    filename: str = "records.csv",
) -> dict[str, Any]:
    r = await client.post(
        f"{_base(grp, zone)}/import/spreadsheet/preview",
        headers=h,
        files={"file": (filename, data, "text/csv")},
        data={"mode": mode},
    )
    assert r.status_code == 200, r.text
    out: dict[str, Any] = r.json()
    return out


async def _commit(
    client: AsyncClient,
    h: dict[str, str],
    grp: DNSServerGroup,
    zone: DNSZone,
    data: bytes,
    *,
    mode: str = "merge",
    confirm: str | None = None,
    filename: str = "records.csv",
    expect: int = 200,
) -> dict[str, Any]:
    form = {"mode": mode}
    if confirm is not None:
        form["confirm_zone"] = confirm
    r = await client.post(
        f"{_base(grp, zone)}/import/spreadsheet/commit",
        headers=h,
        files={"file": (filename, data, "text/csv")},
        data=form,
    )
    assert r.status_code == expect, r.text
    out: dict[str, Any] = r.json()
    return out


async def _seed(db: AsyncSession, zone: DNSZone) -> None:
    db.add_all(
        [
            _rec(zone, "www", "A", "192.0.2.10", ttl=300),
            _rec(zone, "www", "AAAA", "2001:db8::10"),
            _rec(zone, "app", "CNAME", "www"),
            _rec(zone, "@", "MX", "mail.example.com.", priority=10),
            _rec(zone, "@", "TXT", '"v=spf1 mx -all"', tags={"comment": "spf"}),
            _rec(zone, "_sip._tcp", "SRV", "sip.example.com.", priority=10, weight=60, port=5060),
            _rec(zone, "@", "CAA", '0 issue "letsencrypt.org"'),
        ]
    )
    await db.flush()


# ── export ──────────────────────────────────────────────────────────────────


async def test_csv_export_has_one_row_per_record_with_structured_columns(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    await _seed(db_session, zone)
    await db_session.commit()

    r = await client.get(f"{_base(grp, zone)}/export", headers=h, params={"format": "csv"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("text/csv")
    disp = r.headers["content-disposition"]
    assert "dns-zone-sheet.example.edu-" in disp and disp.rstrip('"').endswith(".csv")
    rows = _read_csv(r.content)
    assert list(rows[0].keys()) == COLS
    assert len(rows) == 7
    srv = next(x for x in rows if x["type"] == "SRV")
    assert (srv["priority"], srv["weight"], srv["port"], srv["value"]) == (
        "10",
        "60",
        "5060",
        "sip.example.com.",
    )
    mx = next(x for x in rows if x["type"] == "MX")
    assert (mx["priority"], mx["weight"], mx["port"]) == ("10", "", "")
    caa = next(x for x in rows if x["type"] == "CAA")
    assert caa["value"] == '0 issue "letsencrypt.org"'
    txt = next(x for x in rows if x["type"] == "TXT")
    assert txt["comment"] == "spf"
    a = next(x for x in rows if x["type"] == "A")
    assert a["ttl"] == "300" and uuid.UUID(a["id"])
    aaaa = next(x for x in rows if x["type"] == "AAAA")
    assert aaaa["ttl"] == ""  # inherits the zone's $TTL


async def test_zone_file_export_is_still_the_default(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    await db_session.commit()
    r = await client.get(f"{_base(grp, zone)}/export", headers=h)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/dns")


async def test_export_escapes_formula_injection(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    db_session.add(_rec(zone, "x", "TXT", "=HYPERLINK(1)", tags={"comment": "@SUM(A1)"}))
    await db_session.commit()

    r = await client.get(f"{_base(grp, zone)}/export", headers=h, params={"format": "csv"})
    row = _read_csv(r.content)[0]
    assert row["value"] == "'=HYPERLINK(1)"
    assert row["comment"] == "'@SUM(A1)"

    r = await client.get(f"{_base(grp, zone)}/export", headers=h, params={"format": "xlsx"})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["records"]
    cell = ws.cell(row=2, column=COLS.index("value") + 1)
    assert cell.value == "'=HYPERLINK(1)" and cell.data_type != "f"


# ── round trip ──────────────────────────────────────────────────────────────


async def test_csv_round_trip_is_a_no_op(client: AsyncClient, db_session: AsyncSession) -> None:
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    await _seed(db_session, zone)
    db_session.add(_rec(zone, "f", "TXT", "=1+1"))
    await db_session.commit()

    exported = (
        await client.get(f"{_base(grp, zone)}/export", headers=h, params={"format": "csv"})
    ).content
    p = await _preview(client, h, grp, zone, exported, mode="replace")
    assert p["errors"] == [] and p["skipped"] == []
    assert p["to_create"] == [] and p["to_update"] == [] and p["to_delete"] == []
    assert len(p["unchanged"]) == 8


async def test_xlsx_round_trip_is_a_no_op_and_parses(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    await _seed(db_session, zone)
    await db_session.commit()

    exported = (
        await client.get(f"{_base(grp, zone)}/export", headers=h, params={"format": "xlsx"})
    ).content
    p = await _preview(client, h, grp, zone, exported, mode="replace", filename="r.xlsx")
    assert p["format"] == "xlsx"
    assert p["errors"] == [] and p["to_update"] == [] and p["to_delete"] == []
    assert len(p["unchanged"]) == 7


async def test_xlsx_import_creates_typed_numbers(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    await db_session.commit()
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"  # no ``records`` sheet: the first sheet is read
    ws.append(["Name", "Type", "TTL", "Value", "Priority", "Weight", "Port"])
    ws.append(["_ldap._tcp", "SRV", 600.0, "dc1.example.com.", 0, 100, 389])
    buf = io.BytesIO()
    wb.save(buf)

    body = await _commit(client, h, grp, zone, buf.getvalue(), filename="x.xlsx")
    assert body["created"] == 1
    rec = (
        await db_session.execute(select(DNSRecord).where(DNSRecord.zone_id == zone.id))
    ).scalar_one()
    assert (rec.ttl, rec.priority, rec.weight, rec.port) == (600, 0, 100, 389)


# ── template ────────────────────────────────────────────────────────────────


async def test_template_examples_are_ignored_on_import(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    await db_session.commit()
    r = await client.get(f"{_base(grp, zone)}/import/template", headers=h)
    assert r.status_code == 200
    assert "dns-zone-template-" in r.headers["content-disposition"]
    rows = _read_csv(r.content)
    assert {x["type"] for x in rows} == {"A", "AAAA", "CNAME", "MX", "TXT", "SRV"}
    assert all(x["id"] == "EXAMPLE" for x in rows)

    # The template's examples are valid records, and an operator's row below
    # them imports while the examples do not.
    body = r.content + b",new,A,,192.0.2.99,,,,,,\r\n"
    p = await _preview(client, h, grp, zone, body)
    assert p["examples_ignored"] == 6 and p["errors"] == []
    assert [(c["name"], c["row"]) for c in p["to_create"]] == [("new", 8)]

    # Every example would pass validation if its id were cleared.
    stripped = r.content.replace(b"EXAMPLE", b"")
    p = await _preview(client, h, grp, zone, stripped)
    assert p["errors"] == [] and len(p["to_create"]) == 6

    x = await client.get(
        f"{_base(grp, zone)}/import/template", headers=h, params={"format": "xlsx"}
    )
    assert x.status_code == 200
    assert load_workbook(io.BytesIO(x.content))["records"].max_row == 7


# ── validation + row errors ─────────────────────────────────────────────────


async def test_bad_rows_are_errors_with_row_numbers_and_block_the_commit(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    await db_session.commit()
    data = _csv(
        [
            {"name": "ok", "type": "A", "value": "192.0.2.1"},
            {"name": "bad", "type": "A", "value": "not-an-ip"},
            {"name": "srv", "type": "SRV", "value": "x.example.com.", "priority": 1},
            {"name": "_x", "type": "A", "value": "192.0.2.2"},
            {"name": "t", "type": "A", "value": "192.0.2.3", "ttl": "-5"},
            {"name": "v", "type": "A", "value": "192.0.2.4", "view": "nope"},
            {"name": "w", "type": "BOGUS", "value": "x"},
            {"id": str(uuid.uuid4()), "name": "z", "type": "A", "value": "192.0.2.5"},
        ]
    )
    p = await _preview(client, h, grp, zone, data)
    rows = {e["row"]: e["reason"] for e in p["errors"]}
    assert set(rows) == {3, 4, 5, 6, 7, 8, 9}
    assert "IPv4" in rows[3]
    assert "SRV records require" in rows[4]
    assert "host name" in rows[5]  # BIND check-names (#1378)
    assert "ttl" in rows[6]
    assert "view" in rows[7]
    assert "record_type" in rows[8]
    assert "not a record in this zone" in rows[9]
    assert [c["row"] for c in p["to_create"]] == [2]

    body = await _commit(client, h, grp, zone, data, expect=422)
    assert len(body["detail"]["errors"]) == 7
    count = (await db_session.execute(select(DNSRecord).where(DNSRecord.zone_id == zone.id))).all()
    assert count == []  # never a partial commit


async def test_unreadable_file_is_a_422_not_a_500(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    await db_session.commit()
    for data, name in (
        (b"foo,bar\n1,2\n", "r.csv"),
        (b"not a zip", "r.xlsx"),
        (b"x", "r.json"),
    ):
        r = await client.post(
            f"{_base(grp, zone)}/import/spreadsheet/preview",
            headers=h,
            files={"file": (name, data, "application/octet-stream")},
        )
        assert r.status_code == 422, (name, r.text)


async def test_leading_equals_is_text_on_import(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    await db_session.commit()
    data = _csv(
        [
            {"name": "a", "type": "TXT", "value": "=1+1"},
            {"name": "b", "type": "TXT", "value": "'=2+2"},  # our export's escape
        ]
    )
    await _commit(client, h, grp, zone, data)
    values = {
        r.name: r.value
        for r in (
            await db_session.execute(select(DNSRecord).where(DNSRecord.zone_id == zone.id))
        ).scalars()
    }
    assert values == {"a": "=1+1", "b": "=2+2"}


# ── match rule, merge vs replace ────────────────────────────────────────────


async def test_merge_updates_by_id_and_by_content_and_never_deletes(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    h = await _headers(db_session)
    grp, zone, servers = await _zone(db_session)
    a = _rec(zone, "www", "A", "192.0.2.10", ttl=300)
    mx = _rec(zone, "@", "MX", "mail.example.com.", priority=10)
    keep = _rec(zone, "old", "A", "192.0.2.50")
    db_session.add_all([a, mx, keep])
    await db_session.commit()
    serial_before = zone.last_serial

    data = _csv(
        [
            # by id: rename and re-address
            {"id": str(a.id), "name": "web", "type": "A", "value": "192.0.2.11", "ttl": "300"},
            # by content: priority change
            {"name": "@", "type": "MX", "value": "mail.example.com.", "priority": "20"},
            {"name": "new", "type": "AAAA", "value": "2001:db8::1", "comment": "hello"},
        ]
    )
    p = await _preview(client, h, grp, zone, data)
    assert {u["name"]: u["changed_fields"] for u in p["to_update"]} == {
        "web": ["name", "value"],
        "@": ["priority"],
    }
    assert p["to_delete"] == []

    body = await _commit(client, h, grp, zone, data)
    assert (body["created"], body["updated"], body["deleted"]) == (1, 2, 0)
    await db_session.refresh(a)
    await db_session.refresh(mx)
    await db_session.refresh(zone)
    assert (a.name, a.value, a.fqdn) == ("web", "192.0.2.11", f"web.{ZONE}")
    assert mx.priority == 20
    assert zone.last_serial > serial_before
    new = (await db_session.execute(select(DNSRecord).where(DNSRecord.name == "new"))).scalar_one()
    assert new.tags == {"comment": "hello"}
    assert (await db_session.get(DNSRecord, keep.id)) is not None

    # Agents get record ops: the rename as delete-old + create-new.
    ops = (
        (
            await db_session.execute(
                select(DNSRecordOp).where(DNSRecordOp.server_id == servers[0].id)
            )
        )
        .scalars()
        .all()
    )
    shapes = sorted((o.op, o.record["name"], o.record["value"]) for o in ops)
    assert ("delete", "www", "192.0.2.10") in shapes
    assert ("create", "web", "192.0.2.11") in shapes
    assert ("update", "@", "mail.example.com.") in shapes
    assert ("create", "new", "2001:db8::1") in shapes
    assert {o.target_serial for o in ops} == {zone.last_serial}

    audit = (
        await db_session.execute(
            select(AuditLog).where(AuditLog.resource_type == "dns_zone_import")
        )
    ).scalar_one()
    assert audit.new_value["source"] == "spreadsheet"
    assert audit.new_value["created"] == 1 and audit.new_value["updated"] == 2


async def test_replace_deletes_absent_rows_and_needs_typed_confirmation(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    h = await _headers(db_session)
    grp, zone, servers = await _zone(db_session)
    stay = _rec(zone, "stay", "A", "192.0.2.1")
    gone = _rec(zone, "gone", "A", "192.0.2.2")
    db_session.add_all([stay, gone])
    await db_session.commit()
    data = _csv([{"name": "stay", "type": "A", "value": "192.0.2.1"}])

    p = await _preview(client, h, grp, zone, data, mode="replace")
    assert [d["name"] for d in p["to_delete"]] == ["gone"]

    await _commit(client, h, grp, zone, data, mode="replace", expect=422)
    await _commit(client, h, grp, zone, data, mode="replace", confirm="wrong", expect=422)
    body = await _commit(client, h, grp, zone, data, mode="replace", confirm="sheet.example.edu")
    assert body["deleted"] == 1
    assert (await db_session.get(DNSRecord, gone.id)) is None
    ops = (
        (
            await db_session.execute(
                select(DNSRecordOp).where(DNSRecordOp.server_id == servers[0].id)
            )
        )
        .scalars()
        .all()
    )
    assert [(o.op, o.record["name"]) for o in ops] == [("delete", "gone")]


async def test_comment_only_change_moves_no_serial(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    h = await _headers(db_session)
    grp, zone, servers = await _zone(db_session)
    a = _rec(zone, "a", "A", "192.0.2.1")
    db_session.add(a)
    await db_session.commit()
    serial = zone.last_serial
    data = _csv([{"id": str(a.id), "name": "a", "type": "A", "value": "192.0.2.1", "comment": "c"}])
    body = await _commit(client, h, grp, zone, data)
    assert body["updated"] == 1 and body["target_serial"] is None
    await db_session.refresh(a)
    await db_session.refresh(zone)
    assert a.tags == {"comment": "c"} and zone.last_serial == serial


async def test_type_change_and_double_claim_are_errors(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    a = _rec(zone, "a", "A", "192.0.2.1")
    db_session.add(a)
    await db_session.commit()
    data = _csv(
        [
            {"id": str(a.id), "name": "a", "type": "A", "value": "192.0.2.1"},
            {"name": "a", "type": "A", "value": "192.0.2.1"},
            {"id": str(a.id), "name": "a", "type": "AAAA", "value": "2001:db8::1"},
        ]
    )
    p = await _preview(client, h, grp, zone, data)
    reasons = {e["row"]: e["reason"] for e in p["errors"]}
    assert "row 2 already names this record" in reasons[3]
    assert "type cannot change" in reasons[4]


# ── protected rows, CNAME rule, views ───────────────────────────────────────


async def test_ipam_records_are_never_changed_or_deleted(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    host = _rec(zone, "host", "A", "192.0.2.7", auto_generated=True)
    alias = _rec(zone, "alias", "CNAME", "host", auto_generated=True, ip_alias=True)
    db_session.add_all([host, alias])
    await db_session.commit()

    # Unchanged: reported read-only. Changed: skipped. Absent (replace): kept.
    data = _csv(
        [
            {"id": str(host.id), "name": "host", "type": "A", "value": "192.0.2.7"},
        ]
    )
    p = await _preview(client, h, grp, zone, data, mode="replace")
    assert p["unchanged"][0]["read_only"] is True
    assert p["to_delete"] == []
    assert any("IPAM" in s["reason"] and s["name"] == "alias" for s in p["skipped"])

    data = _csv([{"id": str(host.id), "name": "host", "type": "A", "value": "192.0.2.8"}])
    p = await _preview(client, h, grp, zone, data)
    assert p["to_update"] == []
    assert p["skipped"][0]["row"] == 2 and "IPAM" in p["skipped"][0]["reason"]
    await _commit(client, h, grp, zone, data, mode="replace", confirm=ZONE)
    await db_session.refresh(host)
    assert host.value == "192.0.2.7"
    assert (await db_session.get(DNSRecord, alias.id)) is not None


async def test_cname_rule_clash_is_skipped(client: AsyncClient, db_session: AsyncSession) -> None:
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    db_session.add(_rec(zone, "www", "A", "192.0.2.1"))
    await db_session.commit()
    data = _csv(
        [
            {"name": "www", "type": "CNAME", "value": "elsewhere.example.com."},
            {"name": "ok", "type": "A", "value": "192.0.2.2"},
        ]
    )
    p = await _preview(client, h, grp, zone, data)
    assert [c["name"] for c in p["to_create"]] == ["ok"]
    assert p["skipped"][0]["row"] == 2 and "CNAME" in p["skipped"][0]["reason"]
    body = await _commit(client, h, grp, zone, data)
    assert body["created"] == 1


async def test_view_column_round_trips(client: AsyncClient, db_session: AsyncSession) -> None:
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    inside = DNSView(group_id=grp.id, name="internal")
    db_session.add(inside)
    await db_session.flush()
    db_session.add(_rec(zone, "www", "A", "10.0.0.1", view_id=inside.id))
    db_session.add(_rec(zone, "www", "A", "192.0.2.1"))
    await db_session.commit()

    r = await client.get(f"{_base(grp, zone)}/export", headers=h, params={"format": "csv"})
    rows = _read_csv(r.content)
    assert sorted(x["view"] for x in rows) == ["", "internal"]
    p = await _preview(client, h, grp, zone, r.content, mode="replace")
    assert p["to_update"] == [] and p["to_delete"] == [] and len(p["unchanged"]) == 2

    # Same name/value in a view is a different record from the shared one.
    data = _csv([{"name": "api", "type": "A", "value": "10.0.0.2", "view": "INTERNAL"}])
    await _commit(client, h, grp, zone, data)
    api = (await db_session.execute(select(DNSRecord).where(DNSRecord.name == "api"))).scalar_one()
    assert api.view_id == inside.id


# ── RBAC ────────────────────────────────────────────────────────────────────


async def test_rbac_read_exports_write_imports(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    grp, zone, _ = await _zone(db_session)
    gone = _rec(zone, "gone", "A", "192.0.2.2")
    db_session.add(gone)
    # The export carries the zone and its records, so it needs read on both (#3073).
    reader = await _headers(
        db_session,
        [
            {"action": "read", "resource_type": "dns_zone"},
            {"action": "read", "resource_type": "dns_record"},
        ],
    )
    # The import answers with the zone's records, so it needs read as well.
    writer = await _headers(
        db_session,
        [
            {"action": "write", "resource_type": "dns_record"},
            {"action": "read", "resource_type": "dns_record"},
        ],
    )
    write_only = await _headers(db_session, [{"action": "write", "resource_type": "dns_record"}])
    await db_session.commit()

    r = await client.get(f"{_base(grp, zone)}/export", headers=reader, params={"format": "csv"})
    assert r.status_code == 200
    data = _csv([{"name": "x", "type": "A", "value": "192.0.2.9"}])
    r = await client.post(
        f"{_base(grp, zone)}/import/spreadsheet/commit",
        headers=reader,
        files={"file": ("r.csv", data, "text/csv")},
    )
    assert r.status_code == 403
    await _commit(client, write_only, grp, zone, data, expect=403)
    await _commit(client, writer, grp, zone, data)
    # A replace that deletes also needs delete on dns_record.
    await _commit(client, writer, grp, zone, data, mode="replace", confirm=ZONE, expect=403)


# ── unit: the parser ────────────────────────────────────────────────────────


def test_unescape_only_undoes_the_export_escape() -> None:
    assert sheet_io.unescape_cell("'=x") == "=x"
    assert sheet_io.unescape_cell("'-1") == "-1"
    assert sheet_io.unescape_cell("'plain") == "'plain"
    assert sheet_io.unescape_cell("=x") == "=x"


def test_csv_header_aliases_and_row_numbers() -> None:
    rows = sheet_io.parse_csv(
        b"Name,Record Type,Value\r\nwww,A,192.0.2.1\r\n,,\r\nx,A,192.0.2.2\r\n"
    )
    assert [(r.row, r.name, r.type) for r in rows] == [(2, "www", "A"), (4, "x", "A")]


# ── validator review (#2139) ────────────────────────────────────────────────


async def test_identical_rows_and_an_edit_onto_a_twin_are_errors(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """#1230: the record API refuses to store one RR twice (a 409 on create
    and on an edit); a sheet must not do it either."""
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    a = _rec(zone, "a", "A", "192.0.2.1")
    b = _rec(zone, "b", "A", "192.0.2.2")
    db_session.add_all([a, b])
    await db_session.flush()
    await db_session.commit()
    data = _csv(
        [
            {"name": "new", "type": "A", "value": "192.0.2.9"},
            {"name": "NEW", "type": "A", "value": "192.0.2.9"},
            # An edit by id that turns ``b`` into a copy of ``a``.
            {"id": str(b.id), "name": "a", "type": "A", "value": "192.0.2.1"},
            {"id": str(a.id), "name": "a", "type": "A", "value": "192.0.2.1"},
        ]
    )
    p = await _preview(client, h, grp, zone, data)
    reasons = {e["row"]: e["reason"] for e in p["errors"]}
    assert set(reasons) == {3, 4}, p
    assert "row 2" in reasons[3]
    assert "already holds" in reasons[4] or "row 5" in reasons[4]
    await _commit(client, h, grp, zone, data, expect=422)


async def test_a_legacy_twin_can_still_take_a_ttl_edit(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """Twins that predate #1230 are not re-judged on an edit that keeps the RR."""
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    t1 = _rec(zone, "dup", "A", "192.0.2.1")
    t2 = _rec(zone, "dup", "A", "192.0.2.1")
    db_session.add_all([t1, t2])
    await db_session.flush()
    await db_session.commit()
    data = _csv(
        [
            {"id": str(t1.id), "name": "dup", "type": "A", "value": "192.0.2.1", "ttl": 60},
            {"id": str(t2.id), "name": "dup", "type": "A", "value": "192.0.2.1"},
        ]
    )
    p = await _preview(client, h, grp, zone, data)
    assert p["errors"] == []
    assert [u["existing_id"] for u in p["to_update"]] == [str(t1.id)]


async def test_absolute_names_are_made_relative_or_refused(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """``www.<zone>.`` is ``www``; the record API alone would store it as
    ``www.<zone>`` relative, i.e. ``www.<zone>.<zone>``."""
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    await db_session.commit()
    data = _csv(
        [
            {"name": f"www.{ZONE}", "type": "A", "value": "192.0.2.1"},
            {"name": ZONE.upper(), "type": "TXT", "value": '"apex"'},
            {"name": "host.other.example.", "type": "A", "value": "192.0.2.2"},
            {"name": "rel", "type": "A", "value": "192.0.2.3"},
        ]
    )
    p = await _preview(client, h, grp, zone, data)
    assert [e["row"] for e in p["errors"]] == [4]
    assert "is not in the zone" in p["errors"][0]["reason"]
    assert sorted(c["name"] for c in p["to_create"]) == ["@", "rel", "www"]


async def test_a_legacy_out_of_range_ttl_round_trips(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """A TTL that predates the #2073 bounds is not re-judged when the row
    leaves it as it is; changing it to another out-of-range value is a row
    error, not a 500."""
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    old = _rec(zone, "old", "A", "192.0.2.1", ttl=-5)
    db_session.add(old)
    await db_session.flush()
    await db_session.commit()
    keep = _csv([{"id": str(old.id), "name": "old", "type": "A", "value": "192.0.2.4", "ttl": -5}])
    p = await _preview(client, h, grp, zone, keep)
    assert p["errors"] == [], p
    assert p["to_update"][0]["changed_fields"] == ["value"]
    assert p["to_update"][0]["ttl"] == -5
    bad = _csv([{"id": str(old.id), "name": "old", "type": "A", "value": "192.0.2.1", "ttl": -6}])
    p = await _preview(client, h, grp, zone, bad)
    assert len(p["errors"]) == 1 and "ttl" in p["errors"][0]["reason"]


async def test_the_managed_column_is_never_trusted(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """Clearing ``managed`` on an IPAM record's row does not un-protect it."""
    h = await _headers(db_session)
    grp, zone, _ = await _zone(db_session)
    ipam = _rec(zone, "host", "A", "192.0.2.1", auto_generated=True)
    db_session.add(ipam)
    await db_session.flush()
    await db_session.commit()
    data = _csv(
        [{"id": str(ipam.id), "name": "host", "type": "A", "value": "192.0.2.99", "managed": ""}]
    )
    p = await _preview(client, h, grp, zone, data)
    assert p["to_update"] == []
    assert "IPAM" in p["skipped"][0]["reason"]


def test_csv_utf16_and_bom_are_read() -> None:
    text = "name,type,value\r\nwww,A,192.0.2.1\r\n"
    for data in (
        text.encode("utf-16"),  # BOM + UTF-16LE (Excel's "Unicode text")
        b"\xef\xbb\xbf" + text.encode(),
    ):
        rows = sheet_io.parse_csv(data)
        assert [(r.name, r.value) for r in rows] == [("www", "192.0.2.1")]


def test_blank_rows_are_bounded_before_they_are_all_read() -> None:
    data = b"name,type,value\n" + b",,\n" * (sheet_io.MAX_SHEET_ROWS + 1)
    try:
        sheet_io.parse_csv(data)
    except sheet_io.SpreadsheetError as exc:
        assert "blank ones included" in str(exc)
    else:
        raise AssertionError("a file of blank rows was read in full")


def test_xlsx_that_expands_too_far_is_refused_before_parsing() -> None:
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/sharedStrings.xml", b"\0" * (sheet_io.MAX_XLSX_UNCOMPRESSED_BYTES + 1))
    assert len(buf.getvalue()) < 1024 * 1024  # small on the wire
    try:
        sheet_io.parse_xlsx(buf.getvalue())
    except sheet_io.SpreadsheetError as exc:
        assert "expands to more than" in str(exc)
    else:
        raise AssertionError("an oversized workbook reached openpyxl")


def _xlsx_with_sheet_xml(transform: Any) -> bytes:
    import zipfile

    wb = Workbook()
    ws = wb.active
    ws.title = "records"
    ws.append(["name", "type", "value"])
    ws.append(["www", "A", "192.0.2.1"])
    raw = io.BytesIO()
    wb.save(raw)
    src = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for n in src.namelist():
            d = src.read(n)
            zf.writestr(n, transform(d.decode()).encode() if n == "xl/worksheets/sheet1.xml" else d)
    return out.getvalue()


def test_a_corrupt_or_entity_bomb_sheet_is_a_spreadsheet_error() -> None:
    """read_only parses a sheet lazily, as its rows are iterated: a sheet
    whose XML fails there must be a SpreadsheetError (a 422), not escape as
    a ParseError (a 500)."""
    lol = (
        '<!DOCTYPE lolz [<!ENTITY lol "lol">'
        + "".join(
            f'<!ENTITY lol{i} "' + (f"&lol{i - 1};" if i > 1 else "&lol;") * 10 + '">'
            for i in range(1, 10)
        )
        + "]>"
    )

    def bomb(t: str) -> str:
        # openpyxl writes inline strings and no XML declaration.
        assert "<t>name</t>" in t, t
        return lol + t.replace("<t>name</t>", "<t>&lol9;</t>", 1)

    for data in (
        _xlsx_with_sheet_xml(lambda t: t[: len(t) // 2]),
        _xlsx_with_sheet_xml(bomb),
    ):
        try:
            sheet_io.parse_xlsx(data)
        except sheet_io.SpreadsheetError:
            continue
        raise AssertionError("a broken sheet did not raise SpreadsheetError")
